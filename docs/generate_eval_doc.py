"""Generate a Word report of live NCD gold-set retrieval evaluation results."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCS = Path(__file__).resolve().parent
REPO = DOCS.parent
BACKEND = REPO / "backend"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def set_cell_shading(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def style_table(table) -> None:
    table.style = "Table Grid"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "0F766E")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(9)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_heading_styled(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(15, 23, 42)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = "Calibri"


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = val
    style_table(table)
    doc.add_paragraph()


def _join_ids(ids: list[str] | None) -> str:
    return ", ".join(ids or []) or "—"


def _fmt_rank(rank: int | None) -> str:
    return "—" if rank is None else str(rank)


def run_retrieval() -> tuple[list[dict], dict]:
    from app.eval.runner import load_gold_set, retrieval_metrics
    from app.rag.indexing import index_is_ready

    if not index_is_ready():
        raise FileNotFoundError(
            "Search index not found. Run `python -m app.rag.indexing` from backend/ first."
        )
    cases = load_gold_set()
    metrics = retrieval_metrics(cases, k=5)
    return cases, metrics


def build_docx(cases: list[dict], metrics: dict, out_path: Path) -> None:
    details_by_id = {row["id"]: row for row in metrics["cases"]}
    unanswerable = [c for c in cases if c.get("unanswerable")]
    answerable = [c for c in cases if not c.get("unanswerable") and c.get("expected_doc_ids")]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("NCD Guideline Tutor")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(15, 118, 110)

    sub = doc.add_paragraph()
    r = sub.add_run("Retrieval evaluation results")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(
        doc,
        "This report records a live retrieval evaluation of the WHO NCD Guideline Tutor "
        "against evals/ncd_gold_set.json. Each answerable question is run through "
        "retrieve_contexts with hybrid search (Chroma + BM25 + reciprocal rank fusion) and "
        "FlashRank reranking, top_k = 5. No LLM API key is used. Scores mean “the expected "
        "WHO guideline was retrieved,” not clinical validation.",
    )

    add_heading_styled(doc, "1. Purpose and setup", 1)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Gold set", "evals/ncd_gold_set.json"],
            ["Runner", "app.eval.runner.retrieval_metrics"],
            ["Retrieval", "retrieve_contexts · hybrid=True · rerank=True · k=5"],
            ["Hit definition", "Any top-5 chunk whose doc_id is in expected_doc_ids"],
            ["Answerable cases", str(metrics["n"])],
            ["Unanswerable cases (excluded from scores)", str(len(unanswerable))],
        ],
    )

    add_heading_styled(doc, "2. Headline scores", 1)
    add_body(
        doc,
        "Document hit rate@5 is the fraction of answerable questions where at least one expected "
        "document appears in the top 5 retrieved chunks. MRR is the mean of 1/rank of the "
        "first expected document (0 if none of the expected documents appear at all).",
    )
    add_table(
        doc,
        ["Metric", "Score", "n"],
        [
            ["Document hit rate@5", f"{metrics['document_hit_rate_at_k']:.3f}", str(metrics["n"])],
            ["MRR", f"{metrics['mrr']:.3f}", str(metrics["n"])],
        ],
    )

    add_heading_styled(doc, "3. How scores are computed", 1)
    add_table(
        doc,
        ["Function (eval/metrics.py)", "What it does"],
        [
            [
                "first_relevant_rank",
                "1-based rank of the first retrieved chunk whose doc_id is in expected_doc_ids; None if no match.",
            ],
            [
                "document_hit_rate_at_k",
                "Hits / n, where a hit is rank is not None and rank <= k (here k=5).",
            ],
            [
                "mean_reciprocal_rank",
                "Mean of 1/rank for each case; missing ranks count as 0.",
            ],
        ],
    )
    add_body(
        doc,
        "Unanswerable items (empty expected_doc_ids) are skipped so a correct “not in sources” "
        "pattern is not punished as a retrieval miss.",
    )

    add_heading_styled(doc, "4. Per-case retrieval results", 1)
    add_body(
        doc,
        "Rank is the position of the first expected WHO document in the reranked list. "
        "Retrieved docs are the top-5 doc_id values after hybrid retrieve + rerank.",
    )
    case_rows: list[list[str]] = []
    for case in answerable:
        detail = details_by_id[case["id"]]
        case_rows.append(
            [
                case["id"],
                case["question"],
                _join_ids(case.get("expected_doc_ids")),
                "Yes" if detail["hit_at_k"] else "No",
                _fmt_rank(detail["rank"]),
                _join_ids(detail.get("retrieved_docs")),
            ]
        )
    add_table(
        doc,
        ["Id", "Question", "Expected docs", "Hit@5", "Rank", "Retrieved docs (top 5)"],
        case_rows,
    )

    add_heading_styled(doc, "5. Unanswerable cases (excluded from Document hit rate@5 / MRR)", 1)
    add_body(
        doc,
        "These questions are in the gold set to probe out-of-corpus behaviour. They are not "
        "included in the headline scores.",
    )
    add_table(
        doc,
        ["Id", "Question", "Why excluded"],
        [
            [
                c["id"],
                c["question"],
                "unanswerable=true; no expected_doc_ids",
            ]
            for c in unanswerable
        ],
    )

    add_heading_styled(doc, "6. Caveats", 1)
    add_body(
        doc,
        "These metrics test whether the retriever surfaces the right WHO guideline, not whether "
        "the assistant would give a clinically correct answer. Optional generation proxies "
        "(token overlap, abstention signals, citation source matching) are not included here: "
        "they need an LLM API key and are comparative only. This application is not a medical "
        "device and does not imply WHO endorsement.",
    )

    footer = doc.add_paragraph()
    fr = footer.add_run(
        "WHO source PDFs remain under CC BY-NC-SA 3.0 IGO. Rebuild this report with "
        "`cd backend && PYTHONPATH=. python ../docs/generate_eval_doc.py`."
    )
    fr.italic = True
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(out_path)


def main() -> None:
    out_path = DOCS / "NCD_Guideline_Tutor_Eval_Results.docx"
    cases, metrics = run_retrieval()
    build_docx(cases, metrics, out_path)
    print(
        f"Retrieval n={metrics['n']} Document hit rate@5={metrics['document_hit_rate_at_k']} MRR={metrics['mrr']}"
    )
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
